#!/usr/bin/env python3
"""
Convert a PDF into a PPTX deck: one slide per page, each page rendered
to a high-resolution image via PyMuPDF. Replaces LibreOffice impress_pdf_import.

Usage:
    python3 pdf_to_pptx.py <input.pdf> <output.pptx>
"""
import sys
import os
import tempfile

import fitz
from pptx import Presentation
from pptx.util import Emu


def main():
    if len(sys.argv) < 3:
        sys.stderr.write('Usage: pdf_to_pptx.py <input.pdf> <output.pptx>\n')
        sys.exit(2)
    src, out = sys.argv[1], sys.argv[2]

    doc = fitz.open(src)
    if doc.needs_pass:
        sys.stderr.write('PDF is password-protected.\n')
        sys.exit(1)
    if len(doc) == 0:
        sys.stderr.write('PDF has no pages.\n')
        sys.exit(1)

    first = doc[0]
    prs = Presentation()
    prs.slide_width = Emu(int(first.rect.width * 12700))
    prs.slide_height = Emu(int(first.rect.height * 12700))
    blank = prs.slide_layouts[6]

    tmpdir = tempfile.mkdtemp(prefix='pdf2pptx_')
    try:
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=150)
            img_path = os.path.join(tmpdir, f'page_{i + 1}.png')
            pix.save(img_path)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(
                img_path, 0, 0, width=prs.slide_width, height=prs.slide_height
            )
        prs.save(out)
    finally:
        doc.close()
        for f in os.listdir(tmpdir):
            try:
                os.remove(os.path.join(tmpdir, f))
            except OSError:
                pass
        try:
            os.rmdir(tmpdir)
        except OSError:
            pass

    if not os.path.exists(out) or os.path.getsize(out) == 0:
        sys.stderr.write('Conversion produced no output.\n')
        sys.exit(1)
    sys.exit(0)


if __name__ == '__main__':
    main()
