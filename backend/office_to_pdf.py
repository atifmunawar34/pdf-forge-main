#!/usr/bin/env python3
"""
Convert Office documents (docx, pptx, xlsx) and HTML to PDF
using PyMuPDF's Story engine — no LibreOffice required.

Usage:
    python3 office_to_pdf.py <input.docx|pptx|xlsx|html> <output.pdf>
"""
import sys
import os
import re
import base64
import html as html_mod

import fitz


def render_html_to_pdf(html_text, out_path, mediabox=None):
    """Render an HTML string into a paginated PDF via fitz.Story."""
    story = fitz.Story(html=html_text)
    writer = fitz.DocumentWriter(out_path)
    mb = mediabox or fitz.paper_rect('a4')
    where = mb + (36, 36, -36, -36)
    more = 1
    while more:
        dev = writer.begin_page(mb)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()


def esc(text):
    return html_mod.escape(str(text), quote=False)


def run_to_html(run):
    text = esc(run.text)
    if not text:
        return ''
    if run.bold:
        text = f'<b>{text}</b>'
    if run.italic:
        text = f'<i>{text}</i>'
    if run.underline:
        text = f'<u>{text}</u>'
    try:
        if run.font.size:
            text = f'<span style="font-size:{run.font.size.pt}pt">{text}</span>'
        if run.font.color and run.font.color.rgb:
            text = f'<span style="color:#{run.font.color.rgb}">{text}</span>'
    except Exception:
        pass
    return text


def docx_images(paragraph, doc_part):
    """Yield base64 data URIs for inline images in a paragraph."""
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    for blip in paragraph._element.findall('.//a:blip', ns):
        rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        if not rid or rid not in doc_part.related_parts:
            continue
        part = doc_part.related_parts[rid]
        mime = part.content_type or 'image/png'
        b64 = base64.b64encode(part.blob).decode('ascii')
        yield f'data:{mime};base64,{b64}'


def docx_to_html(path):
    import docx
    doc = docx.Document(path)
    parts = ['<html><body>']
    try:
        blocks = doc.iter_inner_content()
    except AttributeError:
        blocks = list(doc.paragraphs)
    for block in blocks:
        if hasattr(block, 'runs'):  # Paragraph
            style = (block.style.name or '') if block.style else ''
            m = re.match(r'Heading\s+(\d)', style)
            if style == 'Title':
                tag = 'h1'
            elif m:
                tag = f'h{min(int(m.group(1)), 6)}'
            else:
                tag = 'p'
            inner = ''.join(run_to_html(r) for r in block.runs)
            imgs = ''.join(
                f'<img src="{uri}" style="max-width:500px">' for uri in docx_images(block, doc.part)
            )
            if not inner.strip() and not imgs:
                parts.append('<p>&nbsp;</p>')
                continue
            if style == 'List Bullet':
                inner = '&bull; ' + inner
            elif 'List Number' in style:
                inner = '&#8226; ' + inner
            parts.append(f'<{tag}>{inner}</{tag}>{imgs}')
        elif hasattr(block, 'rows'):  # Table
            parts.append('<table border="1" cellspacing="0" cellpadding="4">')
            for row in block.rows:
                parts.append('<tr>')
                for cell in row.cells:
                    cell_html = '<br>'.join(esc(p.text) for p in cell.paragraphs)
                    parts.append(f'<td>{cell_html}</td>')
                parts.append('</tr>')
            parts.append('</table>')
    parts.append('</body></html>')
    return ''.join(parts)


def pptx_to_slide_html(path):
    """Return (slide_html_list, slide_width_pt, slide_height_pt)."""
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation(path)
    w_pt = prs.slide_width / 12700.0
    h_pt = prs.slide_height / 12700.0
    slides = []
    for slide in prs.slides:
        parts = []
        for shape in slide.shapes:
            try:
                if shape.shape_type == 13:  # PICTURE
                    blob = shape.image.blob
                    mime = shape.image.content_type or 'image/png'
                    b64 = base64.b64encode(blob).decode('ascii')
                    parts.append(f'<img src="data:{mime};base64,{b64}" style="max-width:80%">')
                    continue
            except Exception:
                pass
            if shape.has_table:
                parts.append('<table border="1" cellspacing="0" cellpadding="4">')
                for row in shape.table.rows:
                    parts.append('<tr>')
                    for cell in row.cells:
                        parts.append(f'<td>{esc(cell.text)}</td>')
                    parts.append('</tr>')
                parts.append('</table>')
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    inner = ''.join(
                        ('<b>%s</b>' % esc(r.text)) if r.font.bold else esc(r.text)
                        for r in para.runs
                    ) or esc(para.text)
                    if not inner.strip():
                        continue
                    if getattr(shape, 'is_placeholder', False) and shape.placeholder_format is not None:
                        try:
                            if shape.placeholder_format.idx == 0:
                                parts.append(f'<h2>{inner}</h2>')
                                continue
                        except Exception:
                            pass
                    parts.append(f'<p>{inner}</p>')
        slides.append(''.join(parts) or '<p>&nbsp;</p>')
    if not slides:
        slides = ['<p>(empty presentation)</p>']
    return slides, w_pt, h_pt


def xlsx_to_html(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = ['<html><body>']
    for ws in wb.worksheets:
        parts.append(f'<h3>{esc(ws.title)}</h3>')
        parts.append('<table border="1" cellspacing="0" cellpadding="4">')
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if row_count >= 500:
                parts.append('<tr><td><i>(truncated)</i></td></tr>')
                break
            if all(v is None for v in row):
                continue
            parts.append('<tr>')
            for v in row[:30]:
                parts.append(f'<td>{esc("" if v is None else v)}</td>')
            parts.append('</tr>')
            row_count += 1
        parts.append('</table><p>&nbsp;</p>')
    parts.append('</body></html>')
    wb.close()
    return ''.join(parts)


def main():
    if len(sys.argv) < 3:
        sys.stderr.write('Usage: office_to_pdf.py <input> <output.pdf>\n')
        sys.exit(2)
    src, out = sys.argv[1], sys.argv[2]
    ext = os.path.splitext(src)[1].lower()

    if ext in ('.docx', '.docm'):
        render_html_to_pdf(docx_to_html(src), out)
    elif ext in ('.pptx', '.pptm'):
        slides, w_pt, h_pt = pptx_to_slide_html(src)
        writer = fitz.DocumentWriter(out)
        mb = fitz.Rect(0, 0, w_pt, h_pt)
        where = mb + (36, 36, -36, -36)
        for slide_html in slides:
            story = fitz.Story(html=slide_html)
            more = 1
            while more:
                dev = writer.begin_page(mb)
                more, _ = story.place(where)
                story.draw(dev)
                writer.end_page()
        writer.close()
    elif ext in ('.xlsx', '.xlsm', '.xltx', '.xltm'):
        render_html_to_pdf(xlsx_to_html(src), out)
    elif ext in ('.html', '.htm'):
        with open(src, 'r', encoding='utf-8', errors='replace') as f:
            render_html_to_pdf(f.read(), out)
    else:
        sys.stderr.write(f'Unsupported input type: {ext}\n')
        sys.exit(2)

    if not os.path.exists(out) or os.path.getsize(out) == 0:
        sys.stderr.write('Conversion produced no output.\n')
        sys.exit(1)
    sys.exit(0)


if __name__ == '__main__':
    main()
